import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.exceptions import UnsupportedDocument

ALLOWED_SUFFIXES = {".txt", ".pdf", ".docx"}


def normalize_text(text: str, limit: int) -> str:
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:limit]


def extract_resume(filename: str, content: bytes, max_bytes: int, max_chars: int) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise UnsupportedDocument("Поддерживаются только TXT, PDF и DOCX")
    if len(content) > max_bytes:
        raise UnsupportedDocument("Файл превышает допустимый размер")
    try:
        if suffix == ".txt":
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
        else:
            text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
    except Exception as error:
        raise UnsupportedDocument("Не удалось прочитать файл: возможно, он повреждён") from error
    result = normalize_text(text, max_chars)
    if not result:
        raise UnsupportedDocument("В документе нет извлекаемого текста. Для PDF-скана вставьте текст вручную")
    return result

